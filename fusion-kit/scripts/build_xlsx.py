"""Build the companion Excel data dictionary for the P2P/O2C reference document.

Usage: python -I build_xlsx.py <scratchpad_dir> <out.xlsx>
"""
import json
import os
import re
import sys
from collections import Counter

from openpyxl import Workbook
from openpyxl.cell import WriteOnlyCell
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.table import Table, TableStyleInfo

ILLEGAL = re.compile(r"[\x00-\x08\x0b\x0c\x0e-\x1f]")
HEAD_FONT = Font(bold=True, color="FFFFFF")
HEAD_FILL = PatternFill("solid", fgColor="1F4E79")

# chapter filters for P2P / O2C relevance (exact chapter titles from the 26D TOCs, number prefix stripped)
REL_CHAPTERS = {
    "fbdi_fin": {"Budgetary Control", "Cash Management", "Customer Data Model", "Financials Common Module", "General Ledger",
                 "Payables", "Payments", "Receivables", "Revenue Management", "Subledger Accounting", "Tax"},
    "fbdi_proc": {"Purchasing", "Self Service Procurement", "Sourcing", "Supplier Model"},
    "fbdi_scm": {"Inventory Management", "Order Management", "SCM Common"},
    "bicc_fin": {"Advanced Collections", "Budgetary Control", "Cash Management", "Financials Common Module", "General Ledger",
                 "Legal Entity Configurator", "Payables", "Receivables", "Subledger Accounting", "Tax"},
    "bicc_proc": {"Purchasing", "Self Service Procurement", "Sourcing", "Supplier Model", "Supplier Qualification"},
    "bicc_scm": {"Inventory Management", "Order Management", "SCM Common"},
    "ess_proc": {"Self Service Procurement Scheduled Processes", "Supplier Model Scheduled Processes", "Purchasing Scheduled Processes",
                 "Sourcing Scheduled Processes", "Supplier Qualification Management Scheduled Processes"},
    "ess_scm": {"Inventory Management Scheduled Processes", "Order Management Scheduled Processes", "SCM Common Scheduled Processes"},
}
FRONT = {"Previous", "Next", "Title and Copyright Information", "Get Help"}


def chapter(path):
    if not path:
        return ""
    c = path[0]
    return c.split(" ", 1)[1] if c.split(" ", 1)[0].isdigit() and " " in c else c


def clean(v):
    if v is None:
        return ""
    if isinstance(v, (int, float)):
        return v
    v = ILLEGAL.sub(" ", str(v))
    return v[:32000]


def add_sheet(wb, name, headers, rows, widths, table=True):
    ws = wb.create_sheet(title=name[:31])
    ws.freeze_panes = "A2"
    for i, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = w
    head = []
    for h in headers:
        c = WriteOnlyCell(ws, value=h)
        c.font, c.fill = HEAD_FONT, HEAD_FILL
        head.append(c)
    ws.append(head)
    n = 0
    for r in rows:
        ws.append([clean(x) for x in r])
        n += 1
    if n:
        # plain AutoFilter (Excel table objects need column metadata that write-only mode cannot produce safely)
        ws.auto_filter.ref = f"A1:{get_column_letter(len(headers))}{n + 1}"
    return n


def main():
    d, out = sys.argv[1], sys.argv[2]
    cat = json.load(open(os.path.join(d, "catalog.json"), encoding="utf-8"))
    tables = [json.loads(l) for l in open(os.path.join(d, "tables.jsonl"), encoding="utf-8")]
    rest = json.load(open(os.path.join(d, "rest_catalog.json"), encoding="utf-8"))
    toc = json.load(open(os.path.join(d, "guides_toc.json"), encoding="utf-8"))
    p2p_o2c_roots = set()
    rp = os.path.join(d, "doc_sections", "rest_roots_cited.json")
    if os.path.exists(rp):
        p2p_o2c_roots = set(json.load(open(rp, encoding="utf-8")))

    wb = Workbook(write_only=True)
    counts = {}
    flow_order = {"P2P": 0, "O2C": 1, "P2P+O2C": 2, "Shared": 3, "P2P-adjacent": 4, "O2C-adjacent": 5, "Foundation": 6, "Adjacent": 7}
    tables.sort(key=lambda t: (flow_order.get(t["flow"], 9), t["module"], t["name"]))

    readme = [
        ("Oracle Fusion Cloud 26D - P2P and O2C data dictionary (companion to the reference document)",),
        ("Built 2026-10-21 from the official Oracle 'Tables and Views' guides for Financials (oedmf), Procurement (oedmp), SCM (oedsc),",),
        ("Common Features (oedma) and Sales/Fusion Service (oedms), the 'All REST Endpoints' pages of the REST API guides, and the",),
        ("tables of contents of the FBDI, SOAP, OTBI subject area, BICC extract data store, scheduled process and security reference guides.",),
        ("",),
        ("Sheets:",),
        ("  Tables_P2P_O2C  - every documented table in the P2P/O2C/shared modules, with flow tag, PK, column/FK counts, description, doc link",),
        ("  Columns         - every column of those tables (name, datatype, length, precision, not-null, comment, status)",),
        ("  Foreign_Keys    - documented foreign keys (child table -> parent table, child columns)",),
        ("  Indexes         - documented indexes of those tables",),
        ("  All_Tables      - all tables and views in the five guides (16k+ objects), incl. out-of-scope modules and shadow (audit) tables",),
        ("  REST_Operations - every REST operation (method + exact path) in Financials, Procurement, SCM, Common and CX Sales REST APIs",),
        ("  FBDI / BICC / OTBI / SOAP / Sched_Processes / Roles / Business_Objects - TOC entries of those guides with doc links",),
        ("",),
        ("Notes: Fusion SaaS gives no direct database access - tables are reachable through BI Publisher data models (SQL), OTBI,",),
        ("BICC extracts, REST/SOAP and FBDI interface tables. Flow tags are assigned per module (see the document, section 1).",),
        ("Shadow tables (name ending in '_') hold audit/history copies. Always validate against your own pod's release.",),
    ]
    ws = wb.create_sheet("README")
    ws.column_dimensions["A"].width = 150
    for r in readme:
        ws.append(list(r))

    counts["Tables_P2P_O2C"] = add_sheet(
        wb, "Tables_P2P_O2C",
        ["Table", "Flow", "Guide", "Module", "Owner", "Primary key columns", "Columns", "FKs out", "FKs in", "Description", "Oracle doc"],
        ((t["name"], t["flow"], t["guide"], t["module"], t["owner"], "; ".join(p["columns"] for p in t["pk"]), t["n_columns"],
          sum(1 for k in t["fks"] if k["table"].upper() == t["name"]), sum(1 for k in t["fks"] if k["table"].upper() != t["name"]),
          t["description"], t["url"]) for t in tables),
        [34, 12, 7, 30, 8, 34, 8, 8, 8, 100, 60])

    counts["Columns"] = add_sheet(
        wb, "Columns",
        ["Table", "Column", "Datatype", "Length", "Precision", "Not null", "Comment", "Status", "Flow", "Module"],
        ((t["name"], c["name"], c["type"], c["len"], c["prec"], c["nn"], c["comment"], c["status"], t["flow"], t["module"])
         for t in tables for c in t["columns"]),
        [34, 34, 12, 8, 9, 8, 100, 10, 12, 28], table=False)

    edges = sorted({(k["table"].upper(), k["ftable"].upper(), k["col"]) for t in tables for k in t["fks"]})
    inscope = {t["name"] for t in tables}
    counts["Foreign_Keys"] = add_sheet(
        wb, "Foreign_Keys", ["Child table", "Parent table", "Child FK column(s)", "Both in scope"],
        ((c, p, col, "Y" if c in inscope and p in inscope else "N") for c, p, col in edges), [36, 36, 60, 12])

    counts["Indexes"] = add_sheet(
        wb, "Indexes", ["Table", "Index", "Uniqueness", "Columns"],
        ((t["name"], i["name"], i["unique"], i["columns"]) for t in tables for i in t["indexes"]), [34, 40, 14, 80], table=False)

    counts["All_Tables"] = add_sheet(
        wb, "All_Tables", ["Name", "Kind", "Guide", "Module", "Shadow (audit) copy", "In P2P/O2C scope", "Oracle doc"],
        ((r["name"], r["kind"], r["guide"], r["module"], "Y" if r["is_shadow"] else "", "Y" if r["name"] in inscope else "", r["url"])
         for r in sorted(cat, key=lambda r: (r["guide"], r["module"], r["kind"], r["name"]))),
        [40, 8, 7, 40, 10, 10, 70], table=False)

    counts["REST_Operations"] = add_sheet(
        wb, "REST_Operations", ["API book", "Root resource", "Cited in document", "Resource", "Operation", "Method", "Path", "Oracle doc"],
        ((r["book"], r["root"], "Y" if r["root"] in p2p_o2c_roots else "", r["resource"], o["op"], o["method"], o["path"], o["doc"])
         for r in sorted(rest, key=lambda r: (r["book"], r["root"], r["resource"])) for o in r["ops"]),
        [8, 34, 10, 50, 50, 8, 100, 70], table=False)

    def toc_rows(prefixes, relevant_only=True):
        for e in toc:
            if not any(e["book"].startswith(p) for p in prefixes) or e["name"] in FRONT:
                continue
            ch = chapter(e["path"])
            rel = REL_CHAPTERS.get(e["book"])
            flag = "Y" if (rel is None or ch in rel) else ""
            yield (e["book"], ch, " > ".join(e["path"][1:]), e["name"], flag, e["url"])

    hdr = ["Guide", "Chapter", "Sub-path", "Entry", "P2P/O2C relevant chapter", "Oracle doc"]
    w = [12, 34, 40, 70, 12, 80]
    counts["FBDI"] = add_sheet(wb, "FBDI", hdr, toc_rows(["fbdi_"]), w)
    counts["BICC"] = add_sheet(wb, "BICC", hdr, toc_rows(["bicc_"]), w)
    counts["OTBI"] = add_sheet(wb, "OTBI", hdr, toc_rows(["otbi_"]), w)
    counts["SOAP"] = add_sheet(wb, "SOAP", hdr, toc_rows(["soap_"]), w)
    counts["Sched_Processes"] = add_sheet(wb, "Sched_Processes", hdr, toc_rows(["ess_"]), w)
    counts["Roles"] = add_sheet(wb, "Roles", hdr, toc_rows(["sec_"]), w)
    counts["Business_Objects"] = add_sheet(wb, "Business_Objects", hdr, toc_rows(["bo_"]), w)

    wb.save(out)
    print("saved", out, counts)


if __name__ == "__main__":
    main()
