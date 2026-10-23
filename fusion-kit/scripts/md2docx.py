"""Convert the assembled reference Markdown into a styled Word document (python-docx).

Usage: python -I md2docx.py <in.md> <out.docx>
Supports: headings (#..#####), paragraphs, bullet/numbered lists (2 levels), GFM pipe tables, fenced code blocks,
blockquotes, horizontal rules, <!-- pagebreak -->, inline **bold**, *italic*, `code`, [text](url) and bare URLs.
"""
import re
import sys

from docx import Document
from docx.enum.section import WD_ORIENT
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

ACCENT = RGBColor(0x1F, 0x4E, 0x79)
INLINE = re.compile(r"(`[^`]+`|\*\*[^*]+?\*\*|(?<![\w*])\*[^*\s][^*]*?\*(?![\w*])|\[[^\]]+\]\([^)\s]+\)|https?://[^\s)|>]+)")


PPR_AFTER_SHD = {"tabs", "suppressAutoHyphens", "kinsoku", "wordWrap", "overflowPunct", "topLinePunct", "autoSpaceDE",
                 "autoSpaceDN", "bidi", "adjustRightInd", "snapToGrid", "spacing", "ind", "contextualSpacing", "mirrorIndents",
                 "suppressOverlap", "jc", "textDirection", "textAlignment", "textboxTightWrap", "outlineLvl", "divId",
                 "cnfStyle", "rPr", "sectPr", "pPrChange"}
TCPR_AFTER_SHD = {"noWrap", "tcMar", "textDirection", "tcFitText", "vAlign", "hideMark", "headers", "cellIns", "cellDel",
                  "cellMerge", "tcPrChange"}


def shade(cell_or_par, hex_fill):
    is_cell = hasattr(cell_or_par, "_tc")
    pr = cell_or_par._tc.get_or_add_tcPr() if is_cell else cell_or_par._p.get_or_add_pPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), hex_fill)
    after = TCPR_AFTER_SHD if is_cell else PPR_AFTER_SHD
    anchor = next((el for el in pr if el.tag.split("}")[-1] in after), None)
    if anchor is not None:
        anchor.addprevious(shd)
    else:
        pr.append(shd)


def add_field(paragraph, instr):
    r = paragraph.add_run()
    b = OxmlElement("w:fldChar"); b.set(qn("w:fldCharType"), "begin")
    t = OxmlElement("w:instrText"); t.set(qn("xml:space"), "preserve"); t.text = instr
    s = OxmlElement("w:fldChar"); s.set(qn("w:fldCharType"), "separate")
    x = OxmlElement("w:t"); x.text = "Right-click and choose Update Field to build the table of contents." if "TOC" in instr else "1"
    e = OxmlElement("w:fldChar"); e.set(qn("w:fldCharType"), "end")
    for el in (b, t, s):
        r._r.append(el)
    r2 = paragraph.add_run(); r2._r.append(x)
    r3 = paragraph.add_run(); r3._r.append(e)


def add_hyperlink(paragraph, url, text, size=None):
    part = paragraph.part
    rid = part.relate_to(url, "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink", is_external=True)
    h = OxmlElement("w:hyperlink"); h.set(qn("r:id"), rid)
    r = OxmlElement("w:r"); rpr = OxmlElement("w:rPr")
    c = OxmlElement("w:color"); c.set(qn("w:val"), "0563C1"); rpr.append(c)
    u = OxmlElement("w:u"); u.set(qn("w:val"), "single"); rpr.append(u)
    if size:
        sz = OxmlElement("w:sz"); sz.set(qn("w:val"), str(int(size * 2))); rpr.append(sz)
    r.append(rpr)
    t = OxmlElement("w:t"); t.text = text; t.set(qn("xml:space"), "preserve"); r.append(t)
    h.append(r)
    paragraph._p.append(h)


def add_inline(par, text, size=None, bold=False):
    text = text.replace("\\|", "|")
    pos = 0
    for m in INLINE.finditer(text):
        if m.start() > pos:
            run = par.add_run(text[pos:m.start()]); run.bold = bold or None
            if size: run.font.size = Pt(size)
        tok = m.group(0)
        if tok.startswith("`"):
            run = par.add_run(tok[1:-1]); run.font.name = "Consolas"; run.font.size = Pt((size or 10) - 1)
            run._element.rPr.rFonts.set(qn("w:eastAsia"), "Consolas")
        elif tok.startswith("**"):
            run = par.add_run(tok[2:-2]); run.bold = True
            if size: run.font.size = Pt(size)
        elif tok.startswith("*"):
            run = par.add_run(tok[1:-1]); run.italic = True; run.bold = bold or None
            if size: run.font.size = Pt(size)
        elif tok.startswith("["):
            mm = re.match(r"\[([^\]]+)\]\(([^)\s]+)\)", tok)
            if mm.group(2).startswith("http"):
                add_hyperlink(par, mm.group(2), mm.group(1), size)
            else:
                run = par.add_run(mm.group(1))
                if size: run.font.size = Pt(size)
        else:
            add_hyperlink(par, tok, tok, size)
        pos = m.end()
    if pos < len(text):
        run = par.add_run(text[pos:]); run.bold = bold or None
        if size: run.font.size = Pt(size)


def split_row(line):
    line = line.strip()
    if line.startswith("|"):
        line = line[1:]
    if line.endswith("|") and not line.endswith("\\|"):
        line = line[:-1]
    cells, cur, esc = [], "", False
    in_code = False
    for ch in line:
        if ch == "`":
            in_code = not in_code
        if ch == "|" and not esc and not in_code:
            cells.append(cur.strip()); cur = ""
        else:
            cur += ch
        esc = ch == "\\"
    cells.append(cur.strip())
    return cells


def add_table(doc, rows, page_width_cm):
    header, body = rows[0], rows[1:]
    ncol = len(header)
    body = [r + [""] * (ncol - len(r)) if len(r) < ncol else r[:ncol] for r in body]
    big = len(body) > 400
    size = 7.5 if (ncol >= 6 or big) else 8.5
    t = doc.add_table(rows=1 + len(body), cols=ncol)
    t.style = "Table Grid"
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    lens = []
    for j in range(ncol):
        col = [header[j]] + [r[j] for r in body[:300]]
        lens.append(max(4, min(60, max(len(re.sub(r"[`*]", "", c)) for c in col))))
    tot = sum(lens)
    widths = [Cm(max(1.4, page_width_cm * l / tot)) for l in lens]
    for j, h in enumerate(header):
        c = t.rows[0].cells[j]
        c.paragraphs[0].text = ""
        add_inline(c.paragraphs[0], h, size, bold=True)
        for r in c.paragraphs[0].runs:
            r.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
        shade(c, "1F4E79")
    trpr = t.rows[0]._tr.get_or_add_trPr()
    th = OxmlElement("w:tblHeader"); th.set(qn("w:val"), "true"); trpr.append(th)
    for i, r in enumerate(body, 1):
        cells = t.rows[i].cells
        for j, v in enumerate(r):
            p = cells[j].paragraphs[0]
            add_inline(p, v, size)
            p.paragraph_format.space_after = Pt(0)
        if i % 2 == 0 and not big:
            for c in cells:
                shade(c, "EEF3F8")
    for row in t.rows:
        for j, c in enumerate(row.cells):
            c.width = widths[j]
    doc.add_paragraph().paragraph_format.space_after = Pt(2)


def main():
    src, out = sys.argv[1], sys.argv[2]
    lines = open(src, encoding="utf-8").read().splitlines()
    doc = Document()
    sec = doc.sections[0]
    sec.page_width, sec.page_height = Cm(21.0), Cm(29.7)
    sec.orientation = WD_ORIENT.PORTRAIT
    for side in ("left_margin", "right_margin"):
        setattr(sec, side, Cm(1.8))
    sec.top_margin, sec.bottom_margin = Cm(1.8), Cm(1.6)
    page_w = 21.0 - 3.6

    st = doc.styles
    st["Normal"].font.name = "Calibri"; st["Normal"].font.size = Pt(10)
    st["Normal"].paragraph_format.space_after = Pt(4)
    for lvl, sz in ((1, 16), (2, 13), (3, 11.5), (4, 10.5)):
        h = st[f"Heading {lvl}"]; h.font.name = "Calibri"; h.font.size = Pt(sz); h.font.color.rgb = ACCENT
        h.paragraph_format.space_before = Pt(12 if lvl == 1 else 8); h.paragraph_format.space_after = Pt(4)
        h.paragraph_format.keep_with_next = True
    code = st.add_style("CodeBlock", 1)
    code.font.name = "Consolas"; code.font.size = Pt(7.5)
    code.paragraph_format.space_after = Pt(0); code.paragraph_format.space_before = Pt(0)
    code.paragraph_format.left_indent = Cm(0.3)

    # header / footer
    hp = sec.header.paragraphs[0]
    hp.text = "Oracle Fusion Cloud 26D - Procure-to-Pay and Order-to-Cash reference"
    hp.runs[0].font.size = Pt(8); hp.runs[0].font.color.rgb = RGBColor(0x59, 0x59, 0x59)
    fp = sec.footer.paragraphs[0]; fp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    fp.add_run("Page ").font.size = Pt(8); add_field(fp, "PAGE")

    # title page from the first H1 block (lines until first '## ')
    i = 0
    title_lines = []
    while i < len(lines) and not lines[i].startswith("## "):
        title_lines.append(lines[i]); i += 1
    for k, l in enumerate(title_lines):
        l = l.strip()
        if not l:
            continue
        if l.startswith("# "):
            p = doc.add_paragraph(); p.paragraph_format.space_before = Pt(120)
            r = p.add_run(l[2:]); r.bold = True; r.font.size = Pt(26); r.font.color.rgb = ACCENT
        else:
            p = doc.add_paragraph(); add_inline(p, l.lstrip("> "), 11)
    p = doc.add_paragraph(); p.add_run().add_break(WD_BREAK.PAGE)
    p = doc.add_paragraph(); r = p.add_run("Contents"); r.bold = True; r.font.size = Pt(16); r.font.color.rgb = ACCENT
    add_field(doc.add_paragraph(), 'TOC \\o "1-3" \\h \\z \\u')
    p = doc.add_paragraph(); p.add_run().add_break(WD_BREAK.PAGE)

    para, table, in_code, code_lines = [], [], False, []

    def flush_para():
        if para:
            p = doc.add_paragraph(); add_inline(p, " ".join(s.strip() for s in para)); para.clear()

    def flush_table():
        if table:
            rows = [split_row(l) for l in table if not re.match(r"^\s*\|?\s*:?-{2,}", l)]
            if rows:
                add_table(doc, rows, page_w)
            table.clear()

    while i < len(lines):
        line = lines[i]; i += 1
        if in_code:
            if line.strip().startswith("```"):
                p = doc.add_paragraph(style="CodeBlock")
                for k, cl in enumerate(code_lines):
                    if k:
                        p.add_run().add_break()
                    p.add_run(cl.replace("\t", "    "))
                shade(p, "F2F2F2")
                doc.add_paragraph().paragraph_format.space_after = Pt(2)
                in_code, code_lines = False, []
            else:
                code_lines.append(line)
            continue
        if line.strip().startswith("```"):
            flush_para(); flush_table(); in_code = True; continue
        if line.strip().startswith("|"):
            flush_para(); table.append(line); continue
        flush_table()
        s = line.strip()
        if not s:
            flush_para(); continue
        if s == "<!-- pagebreak -->":
            flush_para(); doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE); continue
        if s.startswith("<!--"):
            continue
        m = re.match(r"^(#{1,5})\s+(.*)$", s)
        if m:
            flush_para()
            n = len(m.group(1))
            level = {1: 1, 2: 1, 3: 2, 4: 3, 5: 4}[n]
            doc.add_heading(re.sub(r"[*`]", "", m.group(2)).strip(), level=level)
            continue
        if re.match(r"^(-{3,}|\*{3,}|_{3,})$", s):
            flush_para()
            p = doc.add_paragraph(); ppr = p._p.get_or_add_pPr()
            bdr = OxmlElement("w:pBdr"); b = OxmlElement("w:bottom")
            for k, v in (("w:val", "single"), ("w:sz", "6"), ("w:space", "1"), ("w:color", "A6A6A6")):
                b.set(qn(k), v)
            bdr.append(b); ppr.append(bdr)
            continue
        m = re.match(r"^(\s*)([-*+]|\d+[.)])\s+(.*)$", line)
        if m:
            flush_para()
            indent = len(m.group(1).replace("\t", "    "))
            numbered = m.group(2)[0].isdigit()
            style = ("List Number" if numbered else "List Bullet") + (" 2" if indent >= 2 else "")
            p = doc.add_paragraph(style=style); add_inline(p, m.group(3))
            p.paragraph_format.space_after = Pt(1)
            continue
        if s.startswith(">"):
            flush_para()
            p = doc.add_paragraph(); p.paragraph_format.left_indent = Cm(0.6)
            add_inline(p, s.lstrip("> ").strip()); shade(p, "F3F6FA")
            continue
        para.append(s)
    flush_para(); flush_table()

    # ask Word to refresh fields (TOC, page numbers) when the file opens; keep CT_Settings element order valid
    settings = doc.settings.element
    zoom = settings.find(qn("w:zoom"))
    if zoom is not None and zoom.get(qn("w:percent")) is None:
        zoom.set(qn("w:percent"), "100")
    uf = OxmlElement("w:updateFields"); uf.set(qn("w:val"), "true")
    after_tags = ["hdrShapeDefaults", "footnotePr", "endnotePr", "compat", "docVars", "rsids", "mathPr", "attachedSchema",
                  "themeFontLang", "clrSchemeMapping", "doNotIncludeSubdocsInStats", "doNotAutoCompressPictures", "forceUpgrade",
                  "captions", "readModeInkLockDown", "smartTagType", "schemaLibrary", "shapeDefaults", "doNotEmbedSmartTags",
                  "decimalSymbol", "listSeparator"]
    anchor = next((el for el in settings if el.tag.split("}")[-1] in after_tags), None)
    if anchor is not None:
        anchor.addprevious(uf)
    else:
        settings.append(uf)
    doc.save(out)
    print("saved", out)


if __name__ == "__main__":
    main()
